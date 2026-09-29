#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Preenche a Planilha de Controle de Cartões a partir do relatório
"Relação de Comandas" exportado do FrontOffice.

Uso:
    python preencher_cartoes.py [caminho_do_relatorio]

Se nenhum caminho for passado, o script procura na própria pasta por um
arquivo cujo nome contenha "COMANDA" ou "BORDERO" (.xlsx é preferido; se só
houver .pdf, ele tenta ler o PDF).

O resultado é salvo em PRONTAS/CONTROLE DE CARTOES <data>.xlsx

Regra de classificação de canal (POS/maquininha x ONLINE) - versão atual,
"Opção 1": só olha o que está escrito no campo Documento/Descrição.
    - PIX: sempre POS (maquininha), A NÃO SER que tenha um marcador escrito
      (BEE2PAY, B2B, LINK, etc) no documento - aí conta como ONLINE.
    - Demais bandeiras (cartão): se tiver um marcador escrito, é ONLINE.
      Se NÃO tiver nada escrito no documento, é POS (maquininha).
Não existe mais nenhuma regra baseada no número/prefixo do documento -
essa regra foi abandonada porque o prefixo (NSU do terminal) muda com o
tempo e não é um jeito confiável de identificar a maquininha.
"""
import os
import re
import sys
import glob
import shutil
import zipfile
import unicodedata
import xml.etree.ElementTree as ET
from datetime import datetime, timedelta

import openpyxl
from openpyxl.styles import Font, PatternFill

# ----------------------------------------------------------------------
# CONFIGURAÇÃO
# ----------------------------------------------------------------------
PASTA_BASE = os.path.dirname(os.path.abspath(__file__))
MODELO = os.path.join(PASTA_BASE, "MODELO_CONTROLE_DE_CARTOES.xlsx")
PASTA_SAIDA = os.path.join(PASTA_BASE, "PRONTAS")
NOME_SAIDA = "CONTROLE DE CARTOES {data}.xlsx"  # sem underline, com espaços

LOG = os.path.join(PASTA_BASE, "LOG - quem nao especificou.xlsx")

# marcadores que, escritos no documento/descrição, indicam lançamento ONLINE
MARCADORES_ONLINE = [
    "BEE2PAY", "B2PAY", "BEE 2 PAY", "B2B", "LINK", "BRASPAG", "SITE",
    "PAYMENT",
]

# nomes de operadoras de maquininha que aparecem coladas na Descrição e
# atrapalham a identificação da bandeira (ex.: "Cielo - Elo Credito" tem
# "ELO" dentro de "CIELO")
OPERADORAS = ["CIELO", "REDE", "GETNET", "STONE", "SAFRAPAY", "PAGSEGURO"]

# linha, na aba Controle, de cada bandeira
LINHAS_BANDEIRA = {
    "AMEX": 8,
    "MASTERCARD DÉBITO": 12,
    "MASTERCARD CRÉDITO": 16,
    "VISA DÉBITO": 20,
    "VISA CRÉDITO": 24,
    "ELO CRÉDITO": 28,
    "ELO DÉBITO": 32,
    "PIX": 36,
    "DINERS": 39,
}

RE_LINHA = re.compile(
    r"^(?P<uh>\S+)\s+(?P<tipo>\S+)\s+(?P<reserva>\d+)\s+(?P<conta>\d+)\s+"
    r"(?P<cod>\S+)\s+(?P<descricao>.+?)\s+(?P<valor>-?\d+[.,]\d{2})\s+"
    r"(?P<documento>\S.*?)\s+(?P<data>\d{2}/\d{2}/\d{4})\s+"
    r"(?P<hora>\d{2}:\d{2})\s+(?P<usuario>\S+)\s+(?P<designacao>.*)$"
)
RE_MEIO = re.compile(r"(-?\d+[.,]\d{2})")


# ----------------------------------------------------------------------
# TEXTO / CLASSIFICAÇÃO
# ----------------------------------------------------------------------
def _limpar(txt):
    """Maiúsculas, sem acento, sem nome de operadora, só letras/números/espaço."""
    txt = str(txt or "")
    txt = unicodedata.normalize("NFKD", txt).encode("ascii", "ignore").decode("ascii")
    txt = txt.upper()
    for op in OPERADORAS:
        txt = txt.replace(op, " ")
    txt = re.sub(r"[^A-Z0-9 ]+", " ", txt)
    txt = re.sub(r"\s+", " ", txt).strip()
    return txt


def _marcador_em(texto):
    alvo = _limpar(texto)
    for marca in MARCADORES_ONLINE:
        if re.search(r"\b" + re.escape(_limpar(marca)) + r"\b", alvo):
            return marca
    return None


def classificar_bandeira(descricao, codigo=""):
    d = _limpar(descricao)
    c = _limpar(codigo)
    texto = d + " " + c

    if "PIX" in texto:
        return "PIX"
    if "DINERS" in texto:
        return "DINERS"
    if "AMEX" in texto or "AMERICAN EXPRESS" in texto:
        return "AMEX"
    if "MASTER" in texto:
        if "DEBITO" in texto:
            return "MASTERCARD DÉBITO"
        return "MASTERCARD CRÉDITO"
    if "VISA" in texto:
        if "DEBITO" in texto:
            return "VISA DÉBITO"
        return "VISA CRÉDITO"
    if "ELO" in texto:
        if "DEBITO" in texto:
            return "ELO DÉBITO"
        return "ELO CRÉDITO"
    return "NÃO MAPEADO"


def classificar_canal(documento, descricao=""):
    """
    Opção 1 (regra atual, sem heurística de prefixo de documento):

      - PIX sempre é POS (maquininha), a não ser que tenha um marcador
        (BEE2PAY/B2B/LINK/etc) escrito no documento.
      - Cartão: se tiver marcador escrito, é ONLINE. Se não tiver nada
        escrito no documento, é POS (maquininha).

    Retorna (canal, como_identifiquei).
    """
    doc = str(documento or "").strip()
    marca = _marcador_em(doc + " " + str(descricao or ""))

    if "PIX" in _limpar(descricao):
        if marca:
            return "ONLINE", f"marcador '{marca}'"
        return "POS", "PIX (maquininha)"

    if marca:
        return "ONLINE", f"marcador '{marca}'"

    return "POS", "maquininha (sem marcador no documento)"


# ----------------------------------------------------------------------
# LEITURA DO RELATÓRIO (.xlsx exportado do FrontOffice)
# ----------------------------------------------------------------------
def ler_xlsx(caminho):
    """
    O FrontOffice exporta um .xlsx com o atributo WindowWidth grafado errado
    (deveria ser windowWidth), o que faz o openpyxl recusar o arquivo. Por
    isso lemos o XML na unha (zipfile + ElementTree) em vez de usar openpyxl
    para ler o relatório.

    Colunas (sem cabeçalho, dado começa direto na linha 1):
      A=UH  B=Tipo UH  C=Reserva  D=Conta  E=Cód. Déb.  F=Descrição
      G=Nota (sempre vazia)  H=Valor  I=Documento  J=Data (serial)
      K=Hora (fração do dia)  L=Usuário  M=Designação
    """
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    z = zipfile.ZipFile(caminho)

    strings = []
    if "xl/sharedStrings.xml" in z.namelist():
        root = ET.fromstring(z.read("xl/sharedStrings.xml"))
        for si in root:
            texto = "".join(t.text or "" for t in si.iter(f"{{{ns['m']}}}t"))
            strings.append(texto)

    sheet_name = "xl/worksheets/sheet1.xml"
    sroot = ET.fromstring(z.read(sheet_name))
    sheet_data = sroot.find("m:sheetData", ns)

    linhas = []
    for row in sheet_data:
        cel = {}
        for c in row:
            ref = c.get("r")
            col = "".join(ch for ch in ref if ch.isalpha())
            t = c.get("t")
            v = c.find("m:v", ns)
            val = v.text if v is not None else None
            if t == "s" and val is not None:
                val = strings[int(val)]
            cel[col] = val

        # algumas linhas vêm sem UH preenchida (ex.: reserva já com check-out)
        # mas ainda são lançamentos válidos - o que define se a linha é real
        # é ter um valor, não ter UH.
        valor_txt = cel.get("H")
        if valor_txt in (None, ""):
            continue
        try:
            valor_relatorio = float(valor_txt)
        except ValueError:
            continue
        valor_mais = -valor_relatorio  # inverte o sinal (relatório vem negativo p/ pagamento)

        uh = cel.get("A") or ""

        data_serial = cel.get("J")
        try:
            data_auditoria = (datetime(1899, 12, 30) + timedelta(days=float(data_serial))).date()
        except (TypeError, ValueError):
            data_auditoria = None

        hora_frac = cel.get("K")
        try:
            segundos = round(float(hora_frac) * 24 * 3600)
            hora = f"{segundos // 3600:02d}:{(segundos % 3600) // 60:02d}"
        except (TypeError, ValueError):
            hora = ""

        descricao = cel.get("F") or ""
        documento = cel.get("I") or ""

        linhas.append({
            "uh": uh,
            "tipo_uh": cel.get("B") or "",
            "reserva": cel.get("C") or "",
            "conta": cel.get("D") or "",
            "cod_deb": cel.get("E") or "",
            "descricao": descricao,
            "bandeira": classificar_bandeira(descricao, cel.get("E") or ""),
            "valor_relatorio": valor_relatorio,
            "valor_mais": valor_mais,
            "documento": documento,
            "data": data_auditoria,
            "hora": hora,
            "usuario": cel.get("L") or "",
            "designacao": cel.get("M") or "",
        })

    for linha in linhas:
        canal, como = classificar_canal(linha["documento"], linha["descricao"])
        linha["canal"] = canal
        linha["como_identifiquei"] = como

    return linhas


# ----------------------------------------------------------------------
# LEITURA DO RELATÓRIO (.pdf - Bordero, usado só se não houver .xlsx)
# ----------------------------------------------------------------------
def _texto_do_pdf(caminho):
    try:
        import subprocess
        r = subprocess.run(["pdftotext", "-layout", caminho, "-"],
                            capture_output=True, text=True, timeout=30)
        if r.returncode == 0 and r.stdout.strip():
            return r.stdout
    except Exception:
        pass
    try:
        from pypdf import PdfReader
        reader = PdfReader(caminho)
        return "\n".join(p.extract_text() or "" for p in reader.pages)
    except Exception:
        return ""


def ler_pdf(caminho):
    texto = _texto_do_pdf(caminho)
    linhas = []
    for linha_txt in texto.splitlines():
        linha_txt = linha_txt.strip()
        if not linha_txt:
            continue
        m = RE_LINHA.match(linha_txt)
        if not m:
            continue
        g = m.groupdict()
        try:
            valor_relatorio = float(g["valor"].replace(".", "").replace(",", "."))
        except ValueError:
            continue
        valor_mais = -valor_relatorio
        try:
            data_auditoria = datetime.strptime(g["data"], "%d/%m/%Y").date()
        except ValueError:
            data_auditoria = None

        canal, como = classificar_canal(g["documento"], g["descricao"])
        linhas.append({
            "uh": g["uh"], "tipo_uh": g["tipo"], "reserva": g["reserva"],
            "conta": g["conta"], "cod_deb": g["cod"], "descricao": g["descricao"],
            "bandeira": classificar_bandeira(g["descricao"], g["cod"]),
            "valor_relatorio": valor_relatorio, "valor_mais": valor_mais,
            "documento": g["documento"], "data": data_auditoria, "hora": g["hora"],
            "usuario": g["usuario"], "designacao": g["designacao"],
            "canal": canal, "como_identifiquei": como,
        })
    return linhas


# ----------------------------------------------------------------------
# ACHAR O RELATÓRIO NA PASTA
# ----------------------------------------------------------------------
def achar_relatorio(pasta):
    candidatos = []
    for caminho in glob.glob(os.path.join(pasta, "*")):
        nome = os.path.basename(caminho).upper()
        if "COMANDA" in nome or "BORDERO" in nome:
            candidatos.append(caminho)
    if not candidatos:
        return None
    candidatos.sort(key=lambda c: 0 if c.lower().endswith(".xlsx") else 1)
    return candidatos[0]


# ----------------------------------------------------------------------
# PREENCHER A PLANILHA
# ----------------------------------------------------------------------
def preencher(lancamentos, modelo, destino, data_auditoria):
    wb = openpyxl.load_workbook(modelo)
    ctrl = wb["Controle"]
    cm = wb["Comandas"]

    if data_auditoria:
        ctrl["B3"] = datetime.combine(data_auditoria, datetime.min.time())

    # limpa linhas antigas da aba Comandas (a partir da linha 2)
    if cm.max_row >= 2:
        cm.delete_rows(2, cm.max_row - 1)

    fill_erro = PatternFill("solid", fgColor="FFC7CE")
    fonte_erro = Font(color="9C0006")

    linha_saida = 2
    for lc in sorted(lancamentos, key=lambda x: (x["uh"], x["hora"])):
        cm.cell(row=linha_saida, column=1, value=lc["data"])
        cm.cell(row=linha_saida, column=2, value=lc["uh"])
        cm.cell(row=linha_saida, column=3, value=lc["tipo_uh"])
        cm.cell(row=linha_saida, column=4, value=lc["reserva"])
        cm.cell(row=linha_saida, column=5, value=lc["conta"])
        cm.cell(row=linha_saida, column=6, value=lc["cod_deb"])
        cm.cell(row=linha_saida, column=7, value=lc["descricao"])
        cel_bandeira = cm.cell(row=linha_saida, column=8, value=lc["bandeira"])
        cm.cell(row=linha_saida, column=9, value=lc["canal"])
        cm.cell(row=linha_saida, column=10, value=lc["valor_relatorio"])
        cm.cell(row=linha_saida, column=11, value=lc["valor_mais"])
        cm.cell(row=linha_saida, column=12, value=lc["documento"])
        cm.cell(row=linha_saida, column=13, value=lc["hora"])
        cm.cell(row=linha_saida, column=14, value=lc["usuario"])
        cm.cell(row=linha_saida, column=15, value=lc["designacao"])

        if lc["bandeira"] == "NÃO MAPEADO":
            cel_bandeira.fill = fill_erro
            cel_bandeira.font = fonte_erro

        linha_saida += 1

    wb.security = None  # evita <workbookProtection/> vazio que o openpyxl injeta
    wb.save(destino)


def resumir(lancamentos):
    resumo = {}
    for lc in lancamentos:
        b = resumo.setdefault(lc["bandeira"], {"POS": 0.0, "ONLINE": 0.0, "n": 0})
        b[lc["canal"]] += lc["valor_mais"]
        b["n"] += 1
    return resumo


def brl(v):
    return f"R$ {v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


# ----------------------------------------------------------------------
# LOG PRIVADO - "quem não especificou"
# ----------------------------------------------------------------------
def atualizar_log(pasta, lancamentos):
    """
    Log privado (NÃO é enviado por e-mail, fica só na pasta do auditor) para
    acompanhar quem está lançando sem marcador (BEE2PAY/B2B/etc) no
    documento, quando deveria.

    Importante: com a regra "Opção 1" em vigor, um documento sem marcador
    é simplesmente classificado como POS/maquininha (não gera mais nenhum
    aviso de ambiguidade), porque não dá mais pra saber, só pelo número,
    se aquilo realmente é maquininha ou se era pra ter marcador e a pessoa
    esqueceu de escrever. Ou seja: este log deve, por enquanto, ficar sem
    lançamentos novos - ele fica pronto pra voltar a funcionar assim que
    surgir um jeito confiável de sinalizar isso de novo.
    """
    if os.path.exists(LOG):
        wb = openpyxl.load_workbook(LOG)
    else:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        wb.create_sheet("Lançamentos")
        wb.create_sheet("Resumo")

    lanc_ws = wb["Lançamentos"]
    if lanc_ws.max_row == 1 and lanc_ws["A1"].value is None:
        cabecalho = ["Data", "Usuário", "UH", "Reserva", "Conta", "Bandeira",
                     "Documento", "Valor", "Problema", "Hóspede"]
        for i, titulo in enumerate(cabecalho, start=1):
            c = lanc_ws.cell(row=1, column=i, value=titulo)
            c.font = Font(bold=True)

    existentes = set()
    for row in lanc_ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        existentes.add((str(row[0]), str(row[2]), str(row[3]), str(row[4]),
                         str(row[6]), str(row[7])))

    novos = 0
    for lc in lancamentos:
        problema = None
        como = lc.get("como_identifiquei", "")
        if como.startswith("! SEM MARCADOR") or "! CONFLITO" in como:
            problema = como

        if not problema:
            continue

        chave = (str(lc["data"]), str(lc["uh"]), str(lc["reserva"]),
                 str(lc["conta"]), str(lc["documento"]), str(lc["valor_mais"]))
        if chave in existentes:
            continue
        existentes.add(chave)

        r = lanc_ws.max_row + 1
        lanc_ws.cell(row=r, column=1, value=lc["data"])
        lanc_ws.cell(row=r, column=2, value=lc["usuario"])
        lanc_ws.cell(row=r, column=3, value=lc["uh"])
        lanc_ws.cell(row=r, column=4, value=lc["reserva"])
        lanc_ws.cell(row=r, column=5, value=lc["conta"])
        lanc_ws.cell(row=r, column=6, value=lc["bandeira"])
        lanc_ws.cell(row=r, column=7, value=lc["documento"])
        lanc_ws.cell(row=r, column=8, value=lc["valor_mais"])
        lanc_ws.cell(row=r, column=9, value=problema)
        lanc_ws.cell(row=r, column=10, value=lc["designacao"])
        novos += 1

    # reconstrói a aba Resumo (sempre a primeira aba) com fórmulas vivas
    if "Resumo" in wb.sheetnames:
        del wb["Resumo"]
    resumo_ws = wb.create_sheet("Resumo", 0)
    cabecalho = ["Usuário", "Qtd. sem marcador", "R$ sem marcador",
                 "Qtd. conflitos", "Total de ocorrências"]
    for i, titulo in enumerate(cabecalho, start=1):
        c = resumo_ws.cell(row=1, column=i, value=titulo)
        c.font = Font(bold=True)

    usuarios = sorted({str(row[1]) for row in lanc_ws.iter_rows(min_row=2, values_only=True) if row[1]})
    for i, usuario in enumerate(usuarios, start=2):
        resumo_ws.cell(row=i, column=1, value=usuario)
        resumo_ws.cell(row=i, column=2,
                        value=f'=COUNTIFS(Lançamentos!B:B,A{i},Lançamentos!I:I,"! SEM MARCADOR*")')
        resumo_ws.cell(row=i, column=3,
                        value=f'=SUMIFS(Lançamentos!H:H,Lançamentos!B:B,A{i},Lançamentos!I:I,"! SEM MARCADOR*")')
        resumo_ws.cell(row=i, column=4,
                        value=f'=COUNTIFS(Lançamentos!B:B,A{i},Lançamentos!I:I,"*CONFLITO*")')
        resumo_ws.cell(row=i, column=5, value=f"=B{i}+D{i}")

    wb.save(LOG)
    return novos


# ----------------------------------------------------------------------
# MAIN
# ----------------------------------------------------------------------
def main():
    pasta = PASTA_BASE
    if len(sys.argv) > 1:
        caminho_relatorio = sys.argv[1]
    else:
        caminho_relatorio = achar_relatorio(pasta)

    if not caminho_relatorio or not os.path.exists(caminho_relatorio):
        print("Não encontrei o relatório 'Relação de Comandas' nesta pasta.")
        print("Coloque o arquivo exportado do FrontOffice (.xlsx ou .pdf) aqui")
        print("e rode de novo, ou arraste o arquivo em cima do EXECUTAR.bat.")
        input("\nAperte ENTER para sair...")
        return

    print(f"Lendo: {os.path.basename(caminho_relatorio)}")
    if caminho_relatorio.lower().endswith(".pdf"):
        lancamentos = ler_pdf(caminho_relatorio)
    else:
        lancamentos = ler_xlsx(caminho_relatorio)

    if not lancamentos:
        print("Não consegui ler nenhum lançamento desse relatório.")
        input("\nAperte ENTER para sair...")
        return

    datas = {lc["data"] for lc in lancamentos if lc["data"]}
    data_auditoria = max(datas) if datas else None

    resumo = resumir(lancamentos)
    print(f"\n{len(lancamentos)} lançamentos encontrados"
          + (f" (data: {data_auditoria.strftime('%d/%m/%Y')})" if data_auditoria else "") + "\n")
    print(f"{'Bandeira':<22}{'POS':>16}{'ONLINE':>16}{'Qtd':>6}")
    total_pos = total_online = 0.0
    for bandeira in LINHAS_BANDEIRA:
        d = resumo.get(bandeira, {"POS": 0.0, "ONLINE": 0.0, "n": 0})
        total_pos += d["POS"]
        total_online += d["ONLINE"]
        print(f"{bandeira:<22}{brl(d['POS']):>16}{brl(d['ONLINE']):>16}{d['n']:>6}")
    nao_mapeadas = {b: d for b, d in resumo.items() if b not in LINHAS_BANDEIRA}
    for b, d in nao_mapeadas.items():
        print(f"{b:<22}{brl(d['POS']):>16}{brl(d['ONLINE']):>16}{d['n']:>6}  <-- não mapeada, revisar!")
    print("-" * 60)
    print(f"{'TOTAL':<22}{brl(total_pos):>16}{brl(total_online):>16}{len(lancamentos):>6}")

    sem_marcador = [lc for lc in lancamentos
                    if lc["canal"] == "POS" and "maquininha (sem marcador" in lc["como_identifiquei"]]
    if sem_marcador:
        print(f"\n{len(sem_marcador)} lançamento(s) de cartão sem nenhum marcador escrito no "
              f"documento (BEE2PAY/B2B/LINK/etc) -> classificados como POS/maquininha por padrão.")

    novos_no_log = atualizar_log(pasta, lancamentos)
    if novos_no_log:
        print(f"\n{novos_no_log} novo(s) registro(s) no log privado ({os.path.basename(LOG)}).")

    os.makedirs(PASTA_SAIDA, exist_ok=True)
    if data_auditoria:
        nome_arquivo = NOME_SAIDA.format(data=data_auditoria.strftime("%d-%m-%Y"))
    else:
        nome_arquivo = NOME_SAIDA.format(data=datetime.now().strftime("%d-%m-%Y"))
    destino = os.path.join(PASTA_SAIDA, nome_arquivo)

    if os.path.exists(destino):
        pasta_backup = os.path.join(PASTA_SAIDA, "backup")
        os.makedirs(pasta_backup, exist_ok=True)
        carimbo = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup = os.path.join(pasta_backup, f"{os.path.splitext(nome_arquivo)[0]}__{carimbo}.xlsx")
        shutil.copy2(destino, backup)

    tmp = destino + ".tmp"
    preencher(lancamentos, MODELO, tmp, data_auditoria)
    os.replace(tmp, destino)

    print(f"\nPlanilha pronta: {destino}")
    if nao_mapeadas:
        print("\nATENÇÃO: existem bandeiras não mapeadas, revise a aba Comandas (linhas em vermelho).")

    input("\nAperte ENTER para sair...")


if __name__ == "__main__":
    main()
