"""
Testes da automação de cartões.

Rodar (na pasta do projeto):
    python -m unittest discover -s tests -v
"""
import datetime
import importlib.util
import os
import pathlib
import tempfile
import unittest

import openpyxl

RAIZ = pathlib.Path(__file__).resolve().parents[1]
EXEMPLOS = RAIZ / "exemplos" / "cartoes"


def carregar_modulo():
    spec = importlib.util.spec_from_file_location(
        "preencher_cartoes", RAIZ / "cartoes" / "preencher_cartoes.py")
    modulo = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(modulo)
    return modulo


cartoes = carregar_modulo()


class TestClassificarBandeira(unittest.TestCase):

    def test_bandeiras_basicas(self):
        casos = {
            "Visa Credito": "VISA CRÉDITO",
            "Visa Débito": "VISA DÉBITO",
            "Master Debito": "MASTERCARD DÉBITO",
            "Elo Debito": "ELO DÉBITO",
            "Amex": "AMEX",
            "American Express": "AMEX",
            "Diners": "DINERS",
            "Pix": "PIX",
        }
        for descricao, esperado in casos.items():
            with self.subTest(descricao=descricao):
                self.assertEqual(cartoes.classificar_bandeira(descricao), esperado)

    def test_parcelado_soma_no_credito(self):
        self.assertEqual(cartoes.classificar_bandeira("Mastercard Parcelado"),
                         "MASTERCARD CRÉDITO")

    def test_nome_da_operadora_nao_confunde_a_bandeira(self):
        # "CIELO" tem "ELO" dentro: sem remover a operadora, qualquer bandeira
        # desconhecida passada na Cielo seria contada como Elo
        self.assertEqual(cartoes.classificar_bandeira("Cielo - Hipercard Credito"),
                         "NÃO MAPEADO")
        self.assertEqual(cartoes.classificar_bandeira("Cielo - Elo Credito"), "ELO CRÉDITO")
        self.assertEqual(cartoes.classificar_bandeira("Cielo - Visa Credito"), "VISA CRÉDITO")

    def test_bandeira_desconhecida(self):
        self.assertEqual(cartoes.classificar_bandeira("Hipercard Credito"), "NÃO MAPEADO")


class TestClassificarCanal(unittest.TestCase):

    def canal(self, documento, descricao="Visa Credito"):
        return cartoes.classificar_canal(documento, descricao)[0]

    def test_marcadores_online(self):
        for doc in ("BEE2PAY 48213", "B2B FATURADO", "LINK PGTO 7781", "bee 2 pay"):
            with self.subTest(documento=doc):
                self.assertEqual(self.canal(doc), "ONLINE")

    def test_erro_de_digitacao_b2pay(self):
        self.assertEqual(self.canal("B2PAY"), "ONLINE")

    def test_sem_marcador_e_maquininha(self):
        self.assertEqual(self.canal("383803"), "POS")
        self.assertEqual(self.canal(""), "POS")

    def test_marcador_precisa_ser_palavra_inteira(self):
        # "LINK" dentro de outra palavra não conta como cobrança online
        self.assertEqual(self.canal("LINKADO 123"), "POS")

    def test_pix_e_maquininha_a_menos_que_tenha_marcador(self):
        self.assertEqual(self.canal("849092", "Pix"), "POS")
        self.assertEqual(self.canal("LINK PGTO 1", "Pix"), "ONLINE")


class TestRelatorioDeExemplo(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.lancamentos = cartoes.ler_xlsx(EXEMPLOS / "RELACAO_DE_COMANDAS_DEMO.xlsx")

    def test_le_todos_os_lancamentos(self):
        self.assertEqual(len(self.lancamentos), 22)

    def test_data_da_auditoria(self):
        datas = {lc["data"] for lc in self.lancamentos}
        self.assertEqual(datas, {datetime.date(2026, 9, 28)})

    def test_valores_viram_positivos(self):
        self.assertTrue(all(lc["valor_mais"] > 0 for lc in self.lancamentos))

    def test_separacao_maquininha_online(self):
        online = [lc for lc in self.lancamentos if lc["canal"] == "ONLINE"]
        self.assertEqual(len(online), 6)

    def test_uma_bandeira_nao_mapeada(self):
        nao_mapeadas = [lc for lc in self.lancamentos if lc["bandeira"] == "NÃO MAPEADO"]
        self.assertEqual(len(nao_mapeadas), 1)

    def test_preenche_a_planilha(self):
        with tempfile.TemporaryDirectory() as tmp:
            destino = os.path.join(tmp, "saida.xlsx")
            data = datetime.date(2026, 9, 28)
            cartoes.preencher(self.lancamentos, EXEMPLOS / "MODELO_CONTROLE_DE_CARTOES.xlsx",
                              destino, data)
            wb = openpyxl.load_workbook(destino)
            self.assertEqual(wb["Controle"]["B3"].value.date(), data)
            comandas = wb["Comandas"]
            self.assertEqual(comandas.max_row - 1, 22)
            # bandeira não mapeada fica destacada em vermelho
            vermelhas = [r for r in range(2, comandas.max_row + 1)
                         if comandas.cell(r, 8).value == "NÃO MAPEADO"
                         and comandas.cell(r, 8).fill.fgColor.rgb.endswith("FFC7CE")]
            self.assertEqual(len(vermelhas), 1)


if __name__ == "__main__":
    unittest.main()
